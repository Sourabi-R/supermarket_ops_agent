import os
from decimal import Decimal

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.database.db import get_connection


def generate_invoice_pdf(service, output_path: str = None):
    output = output_path or os.path.join(os.getcwd(), 'sample_outputs', 'sample_invoice.pdf')
    os.makedirs(os.path.dirname(output), exist_ok=True)

    with get_connection() as conn:
        newest_bill = conn.execute(
            "SELECT * FROM bills WHERE status='finalized' ORDER BY finalized_at DESC LIMIT 1"
        ).fetchone()
        if newest_bill:
            bill = dict(newest_bill)
            bill_items = conn.execute(
                "SELECT * FROM bill_items WHERE bill_id = ? ORDER BY id",
                (bill["bill_id"],),
            ).fetchall()
            bill["items"] = [dict(item) for item in bill_items]
        else:
            draft = service.billing.get_open_drafts()
            draft_data = draft[0] if draft else {"draft_id": "DRAFT-DEMO", "customer_name": "Demo Customer", "created_at": "today", "items": []}
            bill = {"bill_id": draft_data.get("draft_id", "DRAFT-DEMO"), "customer_name": draft_data.get("customer_name", "Demo Customer"), "bill_date": draft_data.get("created_at", "today"), "grand_total": "0.00", "payment_mode": "UPI", "items": draft_data.get("items", [])}

    doc = SimpleDocTemplate(output, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Sri Lakshmi Stores", styles['Title']))
    story.append(Paragraph("GSTIN: 29ABCDE1234F1Z5", styles['Normal']))
    story.append(Paragraph(f"Invoice: {bill.get('bill_id', bill.get('draft_id', 'INV-DEMO'))}", styles['Heading2']))
    story.append(Paragraph(f"Customer: {bill.get('customer_name', 'Walk-in')}", styles['Normal']))
    story.append(Paragraph(f"Date: {bill.get('bill_date', bill.get('created_at', 'today'))}", styles['Normal']))
    story.append(Spacer(1, 12))
    data = [["Item", "Qty", "Rate", "GST%", "CGST", "SGST", "Total"]]
    total = Decimal("0")
    for item in bill.get('items', []):
        line = Decimal(str(item.get('line_total', item.get('taxable_amount', '0'))))
        total += line
        data.append([
            item.get('product_name', ''),
            str(item.get('quantity', 0)),
            str(item.get('unit_price', '0')),
            str(item.get('gst_rate_pct', 0)),
            str(item.get('cgst', '0')),
            str(item.get('sgst', '0')),
            str(item.get('line_total', item.get('taxable_amount', '0'))),
        ])
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3A5F')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Grand Total: ₹{float(total):,.2f}", styles['Heading2']))
    doc.build(story)
    return output


def generate_analysis_pptx(service, output_path: str = None):
    output = output_path or os.path.join(os.getcwd(), 'sample_outputs', 'sample_analysis_deck.pptx')
    os.makedirs(os.path.dirname(output), exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = 'Kirana Store Performance'
    title_slide.placeholders[1].text = 'Daily store analytics'

    summary_slide = prs.slides.add_slide(prs.slide_layouts[5])
    summary_slide.shapes.title.text = 'Executive Summary'
    with get_connection() as conn:
        summary = conn.execute("SELECT COUNT(*) AS bill_count, COALESCE(SUM(CAST(grand_total AS REAL)), 0) AS revenue FROM bills WHERE status='finalized'").fetchone()
        top = conn.execute("SELECT product_name, COUNT(*) AS items FROM bill_items GROUP BY product_name ORDER BY items DESC LIMIT 3").fetchall()
    text_box = summary_slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(9), Inches(4.5))
    tf = text_box.text_frame
    tf.text = f"Bills: {summary['bill_count']}\nRevenue: ₹{float(summary['revenue'] or 0):,.2f}\nTop products:\n"
    for row in top:
        tf.text += f"- {row['product_name']}\n"

    chart_slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_slide.shapes.title.text = 'Sales Overview'
    chart_data = CategoryChartData()
    chart_data.categories = ['Revenue', 'Bills']
    chart_data.add_series('Current', (float(summary['revenue'] or 0), int(summary['bill_count'] or 0)))
    chart = chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1.0), Inches(1.6), Inches(10.5), Inches(4.8),
        chart_data
    ).chart
    chart.has_legend = False

    prs.save(output)
    return output
